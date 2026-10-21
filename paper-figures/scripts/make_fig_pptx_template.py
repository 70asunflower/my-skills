# -*- coding: utf-8 -*-
"""Build a paper figure as a NATIVE, fully editable PowerPoint slide
(paper-figures skill). Native rounded-rectangle cards, pill badges, text
boxes and arrow connectors — NOT a pasted image — so the advisor/user can
fine-tune wording and layout in PowerPoint; vector conversion happens
afterwards (see pptx_to_vector.py).

Key conventions (learned by doing):
  * Slide = 2x the final figure size (e.g. figure 3.5 in -> slide 7.0 in);
    fonts and line widths doubled too. LaTeX later scales the exported PDF
    back to the target width — vector is scale-invariant, 2x is comfortable
    to edit.
  * Never change the slide size afterwards: it defines the final crop.
  * Cards = ROUNDED_RECTANGLE with adjustments[0] = radius/min(w,h)
    (pill badges = 0.5). shadow.inherit = False everywhere.
  * Column bodies: ONE textbox per column, paragraphs = lines; bold keyword
    run + regular value run in the same paragraph (word spacing handled by
    the engine — never measure text widths to place a second text object).
  * Line breaks at phrase boundaries; wrapped continuations just flow (PPT
    wraps automatically at the box width).
  * Subscript (V_IT−): run.font._rPr.set("baseline", "-25000").
  * Arrowheads: connectors need a tailEnd appended to a:ln via XML (below).
  * Arrows/lines: shadow.inherit = False or they get theme shadows.
Run:  python make_fig_pptx_template.py   ->  <stem>_edit.pptx next to itself."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

S = 2.0                      # fig-inch -> slide-inch scale (fonts/lines double too)
FIG_W, FIG_H = 3.5, 3.76     # final figure size in inches

def X(i):
    return Inches(i * S)

def Y(i):
    return Inches(i * S)

def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))

FONT = "Arial"
NAVY, NAVY_F, NAVY_T = "1E3D59", "EAF0F6", "0D2538"
RED, RED_F, RED_B, RED_T = "C0392B", "FDEDEC", "C0392B", "78281F"
GRN, GRN_F, GRN_B, GRN_T = "1E8449", "EAFAF1", "27AE60", "145A32"

prs = Presentation()
prs.slide_width = Inches(FIG_W * S)
prs.slide_height = Inches(FIG_H * S)
slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank


def card(x, y, w, h, ec, fc, lw, radius):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), X(w), Y(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fc)
    sh.line.color.rgb = rgb(ec)
    sh.line.width = Pt(lw * S)
    sh.shadow.inherit = False
    return sh


def pill(xc, yc, w, text_, fc, fs=6.2):
    h = 0.17
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                X(xc - w / 2), Y(yc - h / 2), X(w), Y(h))
    sh.adjustments[0] = 0.5                      # stadium pill
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fc)
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text_
    r.font.name = FONT
    r.font.size = Pt(fs * S)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("FFFFFF")


def text(x, y, w, runs, align=PP_ALIGN.LEFT, h=0.26):
    """runs = [(text, fontsize_fig_inch, bold, colorhex)]; y = vertical CENTER."""
    tb = slide.shapes.add_textbox(X(x), Y(y - h / 2), X(w), Y(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    for txt, fs, bold, col in runs:
        r = p.add_run()
        r.text = txt
        r.font.name = FONT
        r.font.size = Pt(fs * S)
        r.font.bold = bold
        r.font.color.rgb = rgb(col)
    return tb


def hline(x1, x2, y, color, lw=0.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x1), Y(y), X(x2), Y(y))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw * S)
    c.shadow.inherit = False


def arrow(x1, y1, x2, y2, color, lw=1.1):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(x1), Y(y1), X(x2), Y(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw * S)
    c.shadow.inherit = False
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "med", "len": "med"}))


# ================= demo: one banner + one card + arrows =================
card(0.125, 0.10, 3.25, 0.52, NAVY, NAVY_F, 1.2, 0.10)
text(0.25, 0.235, 3.0, [("Demo Banner — edit me in PowerPoint", 8.0, True, NAVY_T)],
     align=PP_ALIGN.CENTER)

card(0.125, 0.80, 3.25, 0.70, RED, RED_F, 1.0, 0.10)
pill(1.75, 0.95, 0.68, "[\u2715] EXCLUDED", RED_B)
text(0.245, 1.12, 3.0,
     [("Tier 1: ", 7.2, True, RED_T),
      ("one complete evidence sentence per line, full width, no wrapping", 5.8, False, "333333")])

arrow(1.75, 0.62, 1.75, 0.80, "4A5568")

out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   os.path.basename(__file__).replace("make_", "").replace("_template", "_edit")
                   .replace(".py", ".pptx"))
prs.save(out)
print("saved:", out)
