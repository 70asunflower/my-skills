# -*- coding: utf-8 -*-
"""Markdown -> Word 工程文档转换器（忠实保真，结构化排版）。

用法:
    python convert_doc.py <source.md> <output.docx> [--title "文档标题"] [--no-toc] [--no-header-footer]

设计目标:
    - 把 Markdown 源按"篇/章/节/小节"语义转成规范 Word 排版，而不是简单直译 # 数量。
    - 正文/附录中的引用 [n] / [n,m] / [n-m] 渲染为【上标】（GB/T 7714 顺序编码制）。
    - 文末"参考文献"章节（及其表格）内的 [n] 保持【正体】，不公为上标。
    - 代码段（`...`）内的 [n] 保持原样，不转上标、不转引用。
    - 标题文字中的人工编号（第N章 / 2.1 等）逐字保留，作为唯一编号来源，
      不叠加任何 Word 自动多级编号（避免双重编号）。
    - 100% 内容保真：不增删、不篡改任何文字、数据、术语。

格式默认（可在文件顶部常量区调整）:
    纸张 A4；页边距 上3/下2.5/左2.5/右2.5 cm
    正文 宋体/小四(12pt)/1.5倍行距/首行缩进2字
    标题 黑体；H1=16 H2=14 H3=12 H4=12 H5=10.5；H1-H3 加粗
    代码/英文等宽 Consolas
    目录 TOC 域 TOC \\o "1-3"（仅 H1-H3 进目录，用户打开后需右键"更新域"）
    页眉=文档标题(9pt,下边框)；页脚="第 X 页 / 共 Y 页"(PAGE/NUMPAGES 域)

依赖: python-docx (pip install python-docx)
"""
import re
import argparse
import datetime
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------- 可调整常量 ----------------
CN = "宋体"          # 正文中文字体
HEI = "黑体"         # 标题字体
FANG = "仿宋"        # 引用块/摘字体
EN = "Times New Roman"
CODE = "Consolas"

PAGE_W, PAGE_H = Cm(21), Cm(29.7)
MARGIN_TOP, MARGIN_BOT, MARGIN_L, MARGIN_R = Cm(3), Cm(2.5), Cm(2.5), Cm(2.5)

# 标题字号 / 加粗（按 Word 样式名）
HEADING_SPEC = {
    'Heading 1': (HEI, 16, True),
    'Heading 2': (HEI, 14, True),
    'Heading 3': (HEI, 12, True),
    'Heading 4': (HEI, 12, False),
    'Heading 5': (HEI, 10.5, False),
    'Title': (HEI, 22, True),
    'Quote': (FANG, 12, False),
    'List Bullet': (CN, 12, False),
    'List Number': (CN, 12, False),
    'TOC Heading': (HEI, 16, True),
}

# 强调定界符：开/闭的 * 前若有反斜杠（转义 \*）则不当成定界符，留给 _unescape_md 还原。
INLINE_RE = re.compile(r'(?<!\\)\*\*(.+?)(?<!\\)\*\*|(?<!\\)\*(.+?)(?<!\\)\*|`(.+?)`|\[([^\]]+)\]\(([^)]+)\)')
# 引用序号：纯数字/数字串，如 [1] [1,2] [1-3] [1, 2] [1，2]
CITATION_RE = re.compile(r'\[\d+(?:[\s,，\-]+\d+)*\]')

# 是否处于"文末参考文献"章节（该章节内 [n] 为条目序号，不做上标）
IN_REFERENCES = False
# 当前一级父容器类型：'篇' | '附录' | '前置' | None（用于附录子章节判定）
CUR_PARENT = None


# ---------------- 字体工具 ----------------
def set_run_font(run, cn=CN, en=EN, size=None, bold=None, italic=None, color=None):
    run.font.name = en
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)


def add_superscript_run(paragraph, text, size, bold, italic, cn, en):
    r = paragraph.add_run(text)
    set_run_font(r, cn=cn, en=en, size=size, bold=bold, italic=italic)
    rPr = r._element.get_or_add_rPr()
    va = rPr.find(qn('w:vertAlign'))
    if va is None:
        va = OxmlElement('w:vertAlign')
        rPr.append(va)
    va.set(qn('w:val'), 'superscript')


# Markdown 反斜杠转义还原：去掉"反斜杠 + ASCII 标点"中的反斜杠（CommonMark 行为）。
# 例如 `OC\*`→`OC*`、`RST\(1)`→`RST(1)`、`MR\/CT`→`MR/CT`。
# 仅作用于标点前的转义；反斜杠后接字母（如 Windows 路径 `C:\Users`）保持原样。
_ESCAPE_RE = re.compile(r'\\([!"#$%&\'()*+,\-./:;<=>?@\[\\\]^_`{|}~])')
def _unescape_md(text):
    return _ESCAPE_RE.sub(r'\1', text)


def _add_plain(paragraph, text, size, bold, italic, cn, en):
    if text == '':
        return
    r = paragraph.add_run(_unescape_md(text))
    set_run_font(r, cn=cn, en=en, size=size, bold=bold, italic=italic)


def _add_cit_segments(paragraph, text, size, bold, italic, cn, en):
    """把纯文本段中的引用序号转成上标；参考文献章节内保持正体。"""
    if IN_REFERENCES or not text:
        _add_plain(paragraph, text, size, bold, italic, cn, en)
        return
    pos = 0
    for cm in CITATION_RE.finditer(text):
        if cm.start() > pos:
            _add_plain(paragraph, text[pos:cm.start()], size, bold, italic, cn, en)
        add_superscript_run(paragraph, cm.group(0), size, bold, italic, cn, en)
        pos = cm.end()
    if pos < len(text):
        _add_plain(paragraph, text[pos:], size, bold, italic, cn, en)


def add_runs(paragraph, text, default_size=12, default_bold=False,
             default_italic=False, default_cn=CN, default_en=EN):
    """解析行内格式（加粗/斜体/代码/链接），并把引用序号上标化。"""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _add_cit_segments(paragraph, text[pos:m.start()], default_size,
                              default_bold, default_italic, default_cn, default_en)
        if m.group(1) is not None:
            add_runs(paragraph, m.group(1), default_size, True, default_italic,
                     default_cn, default_en)
        elif m.group(2) is not None:
            add_runs(paragraph, m.group(2), default_size, default_bold, True,
                     default_cn, default_en)
        elif m.group(3) is not None:
            # 代码段保持原样，内部 [n] 不做上标
            r = paragraph.add_run(m.group(3))
            set_run_font(r, cn=CODE, en=CODE, size=default_size, bold=default_bold)
        elif m.group(4) is not None:
            r = paragraph.add_run(m.group(4))
            set_run_font(r, cn=default_cn, en=default_en, size=default_size,
                         bold=default_bold)
            r.font.underline = True
        pos = m.end()
    if pos < len(text):
        _add_cit_segments(paragraph, text[pos:], default_size, default_bold,
                          default_italic, default_cn, default_en)


def set_base_style(style, cn, en, size, bold=None):
    style.font.name = en
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), cn)
    rfonts.set(qn('w:ascii'), en)
    rfonts.set(qn('w:hAnsi'), en)


def style_setup(doc):
    normal = doc.styles['Normal']
    set_base_style(normal, CN, EN, 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(24)  # 中文正文首行缩进2字

    for name, (cn, sz, bold) in HEADING_SPEC.items():
        st = doc.styles[name]
        set_base_style(st, cn, cn if cn in (HEI, FANG) else EN, sz, bold=bold)
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(4)
        if name in ('List Bullet', 'List Number'):
            st.paragraph_format.first_line_indent = Pt(0)
        if name == 'Quote':
            st.paragraph_format.left_indent = Cm(0.75)
            st.paragraph_format.first_line_indent = Pt(0)


# ---------------- 标题层级映射 ----------------
def _heading_style(md_level, text):
    """依据"篇/章/节/小节"语义（而非原文 # 数量）映射 Heading 级别。

    这是本转换器唯一需要按文档微调的地方。常见中文报告/论文约定：
      # 篇 / 附录X / 前置声明        -> Heading 1
      ## 第N章 / 附录直属章节(X.Y)   -> Heading 2
      ### 章内小节(2.1…) / 缺口项(Gx…) -> Heading 3
      #### 措施/模板/证据/明细        -> Heading 4
      ##### 更细明细                  -> Heading 5
    不同文档结构请在此处增删正则与分支。
    """
    global CUR_PARENT
    t = text.strip()
    if md_level <= 1:
        return 'Heading 1'
    if md_level == 2:
        if re.match(r'^第[一二三四五六]篇', t):
            CUR_PARENT = '篇'; return 'Heading 1'
        if re.match(r'^附录\s*[A-EＥ]', t):
            CUR_PARENT = '附录'; return 'Heading 1'
        if t in ('适用对象', '数据可信度声明', '适用边界声明', '适用边界', '全文结构说明'):
            CUR_PARENT = '前置'; return 'Heading 1'
        if re.search(r'第\d+章', t):
            return 'Heading 2'
        return 'Heading 2'
    if md_level == 3:
        if t == '第二篇摘要':
            return 'Heading 2'
        if CUR_PARENT == '篇' and '导读' in t:
            return 'Heading 2'
        if CUR_PARENT == '附录' and re.match(r'^[A-E]\.\d', t):
            return 'Heading 2'
        if re.match(r'^[一二三四五]、', t):
            return 'Heading 2'
        return 'Heading 3'
    if md_level == 4:
        if re.match(r'^\d+\.\d+\.\d+', t):     # 三级子节 14.1.1
            return 'Heading 4'
        if re.match(r'^\d+\.\d+', t):          # 审计项 10.1 / 14.1
            return 'Heading 3'
        if re.match(r'^G\d', t):               # 缺口项 G1-R …
            return 'Heading 3'
        if CUR_PARENT == '附录':
            if t == '概览' or re.match(r'^[一二三四五]、', t) \
                    or t.startswith('主制造商分布') or t.startswith('分类规则说明'):
                return 'Heading 2'
            if t == '缺口总结':
                return 'Heading 3'
        return 'Heading 4'
    return 'Heading 5'


def add_heading(doc, md_level, text):
    global IN_REFERENCES
    style = _heading_style(md_level, text)
    if '参考文献' in text:
        IN_REFERENCES = True
    elif style == 'Heading 1':
        IN_REFERENCES = False
    sizes = {'Heading 1': 16, 'Heading 2': 14, 'Heading 3': 12,
             'Heading 4': 12, 'Heading 5': 10.5}
    p = doc.add_paragraph(style=style)
    p.paragraph_format.first_line_indent = Pt(0)
    add_runs(p, text, default_size=sizes[style],
             default_bold=(style in ('Heading 1', 'Heading 2', 'Heading 3')),
             default_cn=HEI, default_en=HEI)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    add_runs(p, text, default_size=12)
    return p


def add_quote(doc, quote_lines):
    p = doc.add_paragraph(style='Quote')
    for k, ql in enumerate(quote_lines):
        if k > 0:
            p.add_run().add_break()
        add_runs(p, ql, default_size=12, default_cn=FANG, default_en=FANG)
    return p


def add_code_block(doc, code_lines):
    """渲染 ``` 围栏代码块：Consolas 9pt、浅灰底纹 + 细边框（多行连成整体块）、无首行缩进。

    底纹与边框通过 w:shd / w:pBdr 实现；首行带上边框、末行带下边框、
    左右边框每行都有，使多行代码在 Word 中呈现为一个连续的代码块。
    """
    n = len(code_lines)
    for k, cl in enumerate(code_lines):
        p = doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        # 边框：首行 top、末行 bottom、左右恒有（先于其他 pPr 子元素 append，保 schema 顺序）
        pBdr = OxmlElement('w:pBdr')
        for side in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement('w:' + side)
            on = (side != 'top' or k == 0) and (side != 'bottom' or k == n - 1)
            el.set(qn('w:val'), 'single' if on else 'none')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:space'), '3' if side in ('top', 'bottom') else '6')
            el.set(qn('w:color'), 'D1D5DB')
            pBdr.append(el)
        pPr.append(pBdr)
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F3F4F6')
        pPr.append(shd)
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(6) if k == 0 else Pt(0)
        p.paragraph_format.space_after = Pt(6) if k == n - 1 else Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.right_indent = Cm(0.4)
        r = p.add_run(cl if cl else '')
        set_run_font(r, cn=CODE, en=CODE, size=9)
    return None


# ASCII 字符画检测：框线/箭头制表符密集出现的代码块
_ASCII_BOX_CHARS = '─│┌┐└┘├┤┬┴┼▼▲►◄←→═║╔╗╚╝╠╣╦╩╬'


def _looks_like_ascii_art(code_lines):
    """代码块中 ≥3 行、每行含 ≥2 个框线/箭头字符 → 判定为 ASCII 字符画。"""
    hits = 0
    for ln in code_lines:
        if sum(ln.count(ch) for ch in _ASCII_BOX_CHARS) >= 2:
            hits += 1
    return hits >= 3


def add_image(doc, img_path, alt=''):
    """渲染 ![alt](path) 图片：居中、按可用页宽等比缩放、下方附 alt 题注。"""
    import os
    from docx.shared import Cm as _Cm
    if not os.path.isfile(img_path):
        # 图片文件缺失时降级为文字占位，保证不静默丢内容
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        r = p.add_run(f'[图片缺失: {img_path}]')
        set_run_font(r, cn=HEI, en=HEI, size=10)
        return None
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    # A4 可用宽度约 16cm，图片按 15cm 上限等比缩放
    run.add_picture(img_path, width=_Cm(15))
    if alt:
        cap = doc.add_paragraph()
        cap.paragraph_format.first_line_indent = Pt(0)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = cap.add_run(alt)
        set_run_font(rc, cn=HEI, en=HEI, size=10)
    return p


def _fix_bullet_style(doc):
    """把 Word 默认的 List Bullet 改为标准小号项目符号 • (U+2022)，并修正字体。

    Word 内置 List Bullet 的 lvlText 原为 Symbol 字体的私有区字符（如 \\uf0b7），
    直接改字符为 U+2022 但保留 Symbol 字体会渲染为方框（□），因此需同步将
    rFonts 改为标准字体。
    """
    try:
        np_ = doc.part.numbering_part
        nroot = np_.element
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        for absnum in nroot.findall('.//w:abstractNum', ns):
            for lvl in absnum.findall('w:lvl', ns):
                for lt in lvl.findall('w:lvlText', ns):
                    val = lt.get(qn('w:val'), '')
                    # 只改 bullet 类（不含 % 占位符的编号格式）
                    if '%' not in val:
                        lt.set(qn('w:val'), '\u2022')
                        # 同步修复字体：Symbol → 标准字体，否则 • 显示为 □
                        for rPr in lvl.findall('w:rPr', ns):
                            for rf in rPr.findall('w:rFonts', ns):
                                rf.set(qn('w:ascii'), EN)
                                rf.set(qn('w:hAnsi'), EN)
                                rf.set(qn('w:eastAsia'), CN)
    except Exception:
        pass  # 无编号定义时静默跳过


def add_list(doc, items, numbered):
    if numbered:
        # 不依赖 Word 自动编号（否则跨列表续号，且违背 numPr=0 规则）：
        # 直接保留源 markdown 的 "N. " 字面前缀，用悬挂缩进呈现为有序列表，
        # 编号永远与源一一对应，且各列表互不续号。
        for it in items:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.63)
            p.paragraph_format.first_line_indent = Cm(-0.63)
            add_runs(p, it, default_size=12)
    else:
        for it in items:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, it, default_size=12)


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '808080')
    pbdr.append(bottom); pPr.append(pbdr)


def _set_tbl_fixed_layout(t, total_cm):
    """表格固定布局 + 总宽 + 紧凑单元格边距。

    fixed 布局下 gridCol 严格生效：长文本列按设定宽度显示，超长无空格串
    在字符处折行而不是溢出/撑破列宽（autofit 下 Word 会按内容重排，导致
    长列挤占短列、短列被拉宽，可读性差）。
    """
    tblPr = t._tbl.tblPr
    # 复用 python-docx 预置的 <w:tblW w:type="auto" w:w="0"/>；若直接 append 第二个
    # tblW，渲染器会读到排在前面的 w=0，表格退化为"按内容自适应"，列宽全部失效。
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW')
        tblPr.append(w)
    w.set(qn('w:w'), str(int(total_cm * 567)))
    w.set(qn('w:type'), 'dxa')
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', 40), ('left', 80), ('bottom', 40), ('right', 80)):
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def _calc_col_widths(header, data, total_cm=16.0):
    """按每列内容长度分配列宽（cm），短列保底、长列封顶，总和收敛到 total_cm。

    权重取 sqrt(最大字符数) 缓和悬殊比例：长文本列（几百字）不至于吃掉
    全部宽度，短列（位号/日期）不至于窄到换行成竖排。列数越少，单列上限
    越宽（MAXW = total - MINW*(n-1)），保证 2 列"项/要点"表能占满版面。
    """
    n = len(header)
    if n <= 0:
        return []
    MINW = 1.7
    MAXW = max(MINW, total_cm - MINW * (n - 1))   # 其余列保底后本列的最大可用宽
    lens = []
    for j in range(n):
        mx = max([len(header[j])] + [len(r[j]) if j < len(r) else 0 for r in data])
        lens.append(mx ** 0.5 if mx else 1.0)
    s = sum(lens)
    widths = [total_cm * w / s for w in lens]
    # 迭代钳制 [MINW, MAXW] 并把余量摊回未封顶列，使总和 ≈ total_cm
    for _ in range(16):
        widths = [min(max(w, MINW), MAXW) for w in widths]
        slack = total_cm - sum(widths)
        if abs(slack) < 0.03:
            break
        if slack > 0:
            roomy = [j for j in range(n) if widths[j] < MAXW - 0.01]
            if not roomy:
                break
            rs = sum(max(MAXW - widths[j], 0.0) for j in roomy)
            for j in roomy:
                widths[j] += slack * (MAXW - widths[j]) / rs
        else:
            wide = [j for j in range(n) if widths[j] > MINW + 0.01]
            if not wide:
                break
            ws = sum(widths[j] for j in wide)
            for j in wide:
                widths[j] += slack * widths[j] / ws
    widths = [min(max(w, MINW), MAXW) for w in widths]
    return [Cm(w) for w in widths]


def _set_header_repeat(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)


def _set_cell_valign(cell, val='center'):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement('w:vAlign')
    v.set(qn('w:val'), val)
    tcPr.append(v)


def add_table(doc, tbl_lines):
    rows = []
    for tl in tbl_lines:
        if tl.startswith('|'):
            tl = tl[1:]
        if tl.endswith('|'):
            tl = tl[:-1]
        rows.append([c.strip() for c in tl.split('|')])
    header = rows[0]
    sep_idx = None
    for idx, r in enumerate(rows):
        if all(re.match(r'^:?-+:?$', c) for c in r):
            sep_idx = idx
            break
    data = rows[sep_idx + 1:] if sep_idx is not None else rows[1:]
    ncol = max(len(header), max((len(r) for r in data), default=0))
    header = header + [''] * (ncol - len(header))
    t = doc.add_table(rows=1, cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _set_tbl_fixed_layout(t, 16.0)
    col_ws = _calc_col_widths(header, data)
    _set_header_repeat(t.rows[0])
    for j, htext in enumerate(header):
        t.columns[j].width = col_ws[j]
        cell = t.rows[0].cells[j]
        cell.text = ''
        cell.width = col_ws[j]
        cp = cell.paragraphs[0]
        cp.paragraph_format.first_line_indent = Pt(0)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(cp, htext, default_size=10.5, default_bold=True, default_cn=HEI)
        shade_cell(cell, 'D9E2F3')
        _set_cell_valign(cell)
    for r in data:
        r = (r + [''] * (ncol - len(r)))[:ncol]
        cells = t.add_row().cells
        for j, ctext in enumerate(r):
            cell = cells[j]
            cell.text = ''
            cell.width = col_ws[j]
            cp = cell.paragraphs[0]
            cp.paragraph_format.first_line_indent = Pt(0)
            add_runs(cp, ctext, default_size=9)
            _set_cell_valign(cell)


def render_meta_block(doc, mlines, center=False):
    """渲染封面/摘要区的元数据块：正确解析表格与段落（不含标题）。"""
    j, m = 0, len(mlines)
    while j < m:
        s = mlines[j].strip()
        if s == '' or s == '---':
            j += 1
            continue
        if s.startswith('|'):
            tbl = []
            while j < m and mlines[j].strip().startswith('|'):
                tbl.append(mlines[j].strip())
                j += 1
            add_table(doc, tbl)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, s, default_size=12)
        j += 1


def shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def add_field(paragraph, field):
    run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = field
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(instr); run._r.append(f2)
    set_run_font(run, size=9)


def add_toc(doc):
    p = doc.add_paragraph(style='TOC Heading')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('目　录')
    p.paragraph_format.first_line_indent = Pt(0)
    tp = doc.add_paragraph()
    tp.paragraph_format.first_line_indent = Pt(0)
    run = tp.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fsep = OxmlElement('w:fldChar'); fsep.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = '（请在 Word 中右键“更新域”以生成目录）'
    fend = OxmlElement('w:fldChar'); fend.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(instr); run._r.append(fsep)
    run._r.append(t); run._r.append(fend)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def add_header_footer(doc, title):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ''
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(title)
    set_run_font(r, size=9)
    hPr = hp._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '808080')
    pbdr.append(bottom); hPr.append(pbdr)

    fp = sec.footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run('第 '); set_run_font(r1, size=9)
    add_field(fp, 'PAGE')
    r2 = fp.add_run(' 页 / 共 '); set_run_font(r2, size=9)
    add_field(fp, 'NUMPAGES')
    r3 = fp.add_run(' 页'); set_run_font(r3, size=9)


def main():
    ap = argparse.ArgumentParser(description='Markdown -> 规范 Word 文档转换器')
    ap.add_argument('source', help='源 Markdown 文件路径')
    ap.add_argument('output', help='输出 .docx 文件路径')
    ap.add_argument('--title', default=None, help='覆盖文档主标题（默认取第一个 # 标题）')
    ap.add_argument('--no-toc', action='store_true', help='不生成目录页')
    ap.add_argument('--no-header-footer', action='store_true', help='不生成页眉页脚')
    args = ap.parse_args()

    with open(args.source, encoding='utf-8') as f:
        lines = f.read().split('\n')
    doc = Document()
    style_setup(doc)
    sec = doc.sections[0]
    sec.page_width = PAGE_W; sec.page_height = PAGE_H
    sec.top_margin = MARGIN_TOP; sec.bottom_margin = MARGIN_BOT
    sec.left_margin = MARGIN_L; sec.right_margin = MARGIN_R

    # ---- 封面 ----
    title_text = args.title or ''
    i = 0
    n = len(lines)
    while i < n and not lines[i].lstrip().startswith('# '):
        i += 1
    if i < n:
        if not title_text:
            title_text = lines[i].lstrip()[2:].strip()
        i += 1
    # 收集封面元数据（首个 ## 之前）
    meta = []
    while i < n and not lines[i].lstrip().startswith('## '):
        meta.append(lines[i])
        i += 1

    # 分类：作者 / 版本 / 首个属性表 / 其余(摘要·关键词)
    author_line = version_line = None
    attr_tbl, abstract_lines = [], []
    j = 0
    while j < len(meta):
        s = meta[j].strip()
        if s.startswith('**作者'):
            author_line = s; j += 1; continue
        if s.startswith('**版本'):
            version_line = s; j += 1; continue
        if s.startswith('|') and not attr_tbl:
            while j < len(meta) and meta[j].strip().startswith('|'):
                attr_tbl.append(meta[j].strip()); j += 1
            continue
        abstract_lines.append(meta[j]); j += 1

    # 主/副标题拆分（仅封面视觉；页眉/目录仍用完整标题）
    if '——' in title_text:
        main_t, sub_t = title_text.split('——', 1)
    else:
        main_t, sub_t = title_text, ''

    def _blank(k=1):
        for _ in range(k):
            bp = doc.add_paragraph()
            bp.paragraph_format.first_line_indent = Pt(0)
            bp.paragraph_format.space_after = Pt(0)

    _blank(3)
    tp = doc.add_paragraph(style='Title')
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.first_line_indent = Pt(0)
    add_runs(tp, main_t.strip(), default_size=28, default_bold=True,
             default_cn=HEI, default_en=HEI)
    if sub_t.strip():
        stp = doc.add_paragraph()
        stp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stp.paragraph_format.first_line_indent = Pt(0)
        stp.paragraph_format.space_before = Pt(6)
        rs = stp.add_run(sub_t.strip())
        set_run_font(rs, cn=HEI, en=HEI, size=16, bold=False)
    _blank(4)
    # 封面：项目属性表
    if attr_tbl:
        add_table(doc, attr_tbl)
    _blank(3)
    # 封面：作者 / 版本 / 日期（居中）
    now = datetime.datetime.now()
    info = []
    if author_line:
        info.append(re.sub(r'\*\*', '', author_line).strip())
    if version_line:
        vl = re.sub(r'\*\*', '', version_line).strip()
        # 版本行按 "｜" 拆成多行居中显示：避免一整条超长串（含作者/日期/性质）换行不美观
        info.extend(p.strip() for p in vl.split('｜') if p.strip())
    # 版本行已含"日期"字样时不再自动补日期行，避免与版本行日期重复
    if not any('日期' in it for it in info):
        info.append('日期：%d 年 %d 月' % (now.year, now.month))
    for it in info:
        ip = doc.add_paragraph()
        ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ip.paragraph_format.first_line_indent = Pt(0)
        ip.paragraph_format.space_after = Pt(3)
        ri = ip.add_run(it)
        set_run_font(ri, cn=CN, en=EN, size=13)
    add_page_break(doc)

    # ---- 摘要页（正确解析表格/段落）----
    render_meta_block(doc, abstract_lines)
    add_page_break(doc)
    if not args.no_toc:
        add_toc(doc)
        add_page_break(doc)

    # ---- 正文 ----
    bullet_re = re.compile(r'^-\s+')
    num_re = re.compile(r'^\d+\.\s+')
    while i < n:
        line = lines[i]
        s = line.strip()
        if s == '':
            i += 1; continue
        if s == '---':
            add_hr(doc); i += 1; continue
        img_m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', s)
        if img_m:
            add_image(doc, img_m.group(2).strip(), img_m.group(1).strip())
            i += 1; continue
        if s.startswith('```'):
            code = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            i += 1  # 跳过结束围栏
            if _looks_like_ascii_art(code):
                print('[提示] 检测到疑似 ASCII 字符画（框线图/拓扑图/链路图）。'
                      '建议先用 fireworks-tech-graph skill 重绘为正式 SVG/PNG，'
                      '再以 ![题注](图片.png) 替换该代码块并重新转换——'
                      'ASCII 字符画直接转 docx 只能得到等宽文本块，观感差且不正式。当前按等宽文本块渲染。')
            add_code_block(doc, code); continue
        if s.startswith('#'):
            level = len(s) - len(s.lstrip('#'))
            text = s[level:].strip()
            add_heading(doc, level, text); i += 1; continue
        if s.startswith('|'):
            tbl = []
            while i < n and lines[i].strip().startswith('|'):
                tbl.append(lines[i].strip()); i += 1
            add_table(doc, tbl); continue
        if s.startswith('>'):
            ql = []
            while i < n and lines[i].strip().startswith('>'):
                ql.append(re.sub(r'^>\s?', '', lines[i].strip())); i += 1
            add_quote(doc, ql); continue
        if bullet_re.match(s):
            items = []
            while i < n and bullet_re.match(lines[i].strip()):
                items.append(bullet_re.sub('', lines[i].strip())); i += 1
            add_list(doc, items, numbered=False); continue
        if num_re.match(s):
            items = []
            while i < n and num_re.match(lines[i].strip()):
                items.append(lines[i].strip())  # 保留 "N. " 字面编号，避免续号
                i += 1
            add_list(doc, items, numbered=True); continue
        add_paragraph(doc, s); i += 1

    if not args.no_header_footer:
        add_header_footer(doc, title_text)
    doc.core_properties.title = title_text
    _fix_bullet_style(doc)
    doc.save(args.output)

    hcount = {1: 0, 2: 0, 3: 0, 4: 0, 5: 0}
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            try:
                hcount[int(p.style.name.split()[-1])] += 1
            except ValueError:
                pass
    print('OK ->', args.output)
    print('标题统计 H1/H2/H3/H4/H5 =', hcount)
    print('表格数 =', len(doc.tables))


if __name__ == '__main__':
    main()
