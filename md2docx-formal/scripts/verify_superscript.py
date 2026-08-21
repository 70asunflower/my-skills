# -*- coding: utf-8 -*-
"""核验 docx 中引用序号 [n] 的上标处理是否符合 GB/T 7714 顺序编码制。

判定规则：
    - 正文 / 附录（参考文献章节之外）的 [n] 应为【上标】。
    - 参考文献章节（标题含"参考文献"的 Heading，到下一个 Heading 1 为止）内的
      [n] 为条目序号，应为【正体】。
    - 表格中首格="编号"的参考文献表，[1]-[69] 等应为【正体】。

用法:  python verify_superscript.py <file.docx>
退出码: 0=通过, 1=存在违规（并打印违规位置）。
依赖: python-docx
"""
import sys
import re
from docx import Document
from docx.oxml.ns import qn

CIT = re.compile(r'\[\d+(?:[\s,，\-]+\d+)*\]')


def is_superscript(run):
    rp = run._element.find(qn('w:rPr'))
    if rp is None:
        return False
    va = rp.find(qn('w:vertAlign'))
    return va is not None and va.get(qn('w:val')) == 'superscript'


def main():
    if len(sys.argv) < 2:
        print('用法: python verify_superscript.py <file.docx>')
        sys.exit(2)
    path = sys.argv[1]
    doc = Document(path)
    paras = doc.paragraphs

    # 1) 定位参考文献章节区间 [ref_start, ref_end)
    ref_start = None
    for idx, p in enumerate(paras):
        if p.style.name.startswith('Heading') and '参考文献' in p.text:
            ref_start = idx
            break
    ref_end = len(paras)
    if ref_start is not None:
        for idx in range(ref_start + 1, len(paras)):
            if paras[idx].style.name == 'Heading 1':
                ref_end = idx
                break

    viol_in_ref = []   # 章节内被误标上标（应0）
    viol_out = 0       # 章节外未上标（应0）
    body_sup = 0
    for idx, p in enumerate(paras):
        in_ref = (ref_start is not None and ref_start <= idx < ref_end)
        for r in p.runs:
            if not CIT.search(r.text):
                continue
            sup = is_superscript(r)
            if in_ref and sup:
                viol_in_ref.append((idx, p.text[:34], r.text[:12]))
            elif (not in_ref) and not sup:
                viol_out += 1
            elif (not in_ref) and sup:
                body_sup += 1

    # 2) 参考文献表（位于参考文献章节区间内、且首格="编号"）必须全平体
    #    注意：附录等其它表也可能以"编号"为首列，故必须限定在参考文献章节区间内，
    #          避免把附录里本应上标的引用误判为违规（例如 附录D 的 [73]-[76]）。
    table_viol = []
    if ref_start is not None:
        body = doc.element.body
        para_index = {p._p: i for i, p in enumerate(paras)}
        blocks = list(body)

        def table_block_index(tbl):
            last_p = None
            for el in blocks:
                if el.tag == qn('w:p'):
                    if el in para_index:
                        last_p = para_index[el]
                elif el is tbl._tbl:
                    return last_p
            return None

        for t in doc.tables:
            bi = table_block_index(t)
            if bi is None or not (ref_start <= bi < ref_end):
                continue
            if not (t.rows and t.rows[0].cells[0].text.strip() == '编号'):
                continue
            for ri, row in enumerate(t.rows):
                for c in row.cells:
                    if CIT.search(c.text) and any(is_superscript(r)
                                                  for r in c.paragraphs[0].runs):
                        table_viol.append((ri, c.text[:14]))

    print('参考文献章节区间: [%s, %s)' % (ref_start, ref_end))
    print('正文/附录 引用上标数 =', body_sup)
    print('参考文献章节内 误标上标(应0) =', len(viol_in_ref))
    for v in viol_in_ref[:20]:
        print('   ', v)
    print('章节外 未上标(应0) =', viol_out)
    print('参考文献表 上标违规(应0) =', len(table_viol))
    for v in table_viol[:20]:
        print('   ', v)

    ok = (len(viol_in_ref) == 0 and viol_out == 0 and len(table_viol) == 0)
    print('\n结论:', '✅ 通过' if ok else '⚠ 存在违规')
    sys.exit(0 if ok else 1)


if __name__ == '__main__':
    main()
