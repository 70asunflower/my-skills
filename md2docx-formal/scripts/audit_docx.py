# -*- coding: utf-8 -*-
"""对生成的 docx 做结构化格式审计，输出可定位的问题清单。

检查项：
    A. 标题层级 + Word 自动编号（双重编号风险）
    B. 各级标题字体/字号是否统一
    C. 正文(Normal)字体/字号
    D. 目录(TOC)域 / 页眉 / 页脚 / 页码 / 页面尺寸
    E. 表格数量
    F. （可选）与源 MD 保真度：标题数对比 + 关键术语/数据是否健在

用法:
    python audit_docx.py <file.docx> [--source source.md] [--checks 词1,词2,...]
依赖: python-docx
"""
import sys
import re
import argparse
from docx import Document
from docx.oxml.ns import qn


def has_numpr(p):
    pPr = p._p.find(qn('w:pPr'))
    return pPr is not None and pPr.find(qn('w:numPr')) is not None


def main():
    ap = argparse.ArgumentParser(description='docx 格式审计')
    ap.add_argument('docx', help='待审计的 .docx 文件')
    ap.add_argument('--source', default=None, help='源 Markdown（用于保真度对比）')
    ap.add_argument('--checks', default='', help='保真度必含关键词，逗号分隔')
    args = ap.parse_args()

    doc = Document(args.docx)
    sec = doc.sections[0]
    out = []

    # A. 标题层级 / 自动编号
    out.append('=' * 70)
    out.append('A. 标题层级 / 自动编号（双重编号检测）')
    out.append('=' * 70)
    levels, double_num = {}, []
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            levels[p.style.name] = levels.get(p.style.name, 0) + 1
            if has_numpr(p):
                double_num.append(p.text)
    for lvl in sorted(levels, key=lambda x: int(x.replace('Heading ', ''))):
        out.append(f'  {lvl:<10}: {levels[lvl]} 个')
    msg = f'  ★ 标题带 Word 自动编号(numPr)数量 = {len(double_num)}'
    out.append(msg + ('  -> ⚠ 存在双重编号风险' if double_num else '  -> ✓ 无'))
    for t in double_num[:10]:
        out.append(f'      └ {t[:50]}')

    # B. 标题字体/字号
    out.append('\n' + '=' * 70)
    out.append('B. 各级标题实际字体 / 字号')
    out.append('=' * 70)
    from collections import defaultdict
    font_map, size_map = defaultdict(set), defaultdict(set)
    for p in doc.paragraphs:
        if p.style.name.startswith('Heading'):
            for r in p.runs:
                if r.font.name:
                    font_map[p.style.name].add(r.font.name)
                if r.font.size:
                    size_map[p.style.name].add(round(r.font.size.pt, 1))
    for lvl in sorted(font_map, key=lambda x: int(x.replace('Heading ', ''))):
        out.append(f'  {lvl:<10} 字体={sorted(font_map[lvl])}  字号={sorted(size_map[lvl])}')

    # C. 正文
    out.append('\n' + '=' * 70)
    out.append('C. 正文(Normal)字体 / 字号')
    out.append('=' * 70)
    nf, ns = set(), set()
    for p in doc.paragraphs:
        if p.style.name == 'Normal':
            for r in p.runs:
                if r.font.name:
                    nf.add(r.font.name)
                if r.font.size:
                    ns.add(round(r.font.size.pt, 1))
    out.append(f'  正文 run 字体集合 = {sorted(nf)}')
    out.append(f'  正文 run 字号集合 = {sorted(ns)}')

    # D. TOC/页眉/页脚/页码/页面
    out.append('\n' + '=' * 70)
    out.append('D. 目录(TOC) / 页眉 / 页脚 / 页码 / 页面尺寸')
    out.append('=' * 70)
    toc_xml = any('TOC' in p._p.xml for p in doc.paragraphs)
    hdr = sec.header.paragraphs[0].text
    ftr_xml = sec.footer.paragraphs[0]._p.xml
    out.append(f'  TOC 域存在 = {toc_xml}')
    out.append(f'  页眉文字 = {repr(hdr[:40])}')
    out.append(f'  页脚含 PAGE 域 = {"PAGE" in ftr_xml}; 含 NUMPAGES 域 = {"NUMPAGES" in ftr_xml}')
    out.append(f'  页面尺寸 = {sec.page_width.cm:.1f} x {sec.page_height.cm:.1f} cm')

    # E. 表格
    out.append('\n' + '=' * 70)
    out.append('E. 表格数量')
    out.append('=' * 70)
    out.append(f'  表格数 = {len(doc.tables)}')

    # F. 保真度
    if args.source or args.checks:
        out.append('\n' + '=' * 70)
        out.append('F. 与源 MD 保真度')
        out.append('=' * 70)
        blocks = [p.text for p in doc.paragraphs] + \
                 [c.text for t in doc.tables for r in t.rows for c in r.cells]
        full = '\n'.join(blocks)
        if args.checks:
            checks = [c.strip() for c in args.checks.split(',') if c.strip()]
            miss = [c for c in checks if c not in full]
            out.append(f'  关键数据/术语缺失 = {len(miss)}' + (str(miss) if miss else '  ✓ 全部存在'))
        if args.source:
            src_lines = open(args.source, encoding='utf-8').read().splitlines()
            src_h = sum(1 for l in src_lines if re.match(r'^#{1,5} ', l))
            doc_h = sum(1 for p in doc.paragraphs if p.style.name.startswith('Heading'))
            out.append(f'  源标题(1-5级)数 = {src_h}; docx 标题数 = {doc_h}  (差={src_h - doc_h})')

    print('\n'.join(out))


if __name__ == '__main__':
    main()
