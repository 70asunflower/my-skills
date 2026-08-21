# -*- coding: utf-8 -*-
"""源 Markdown ↔ 生成 docx 的【内容一致性审查】（逐块保真核验）。

与 audit_docx.py 的 F 段（只看标题数 + 硬编码关键词）不同，本脚本做真正的
逐块覆盖比对，能定位到"哪一句/哪个单元格"没转进去或被改动：

    正向覆盖（源 -> docx，主判定）：
        把源里每个文本单元（标题/段落/表格单元格/引用块/列表项）按转换器的
        搬运规则规范化后，检查它是否作为子串出现在 docx 全文中。任一缺失即判失败。
    反向抽查（docx -> 源，信息性）：
        把 docx 正文里较长的段落回查源文本，列出疑似"新增/编造"的内容，
        并用白名单排除转换器自加的封面/目录/页眉页脚模板（日期、目录占位等）。

文本规范化严格对齐 convert_doc.py 的搬运方式：
    - 链接 [text](url) 只保留 text（url 被转换器丢弃）
    - 去掉 **加粗** / *斜体* / `代码` 标记，保留内部文字
    - 表格按单元格拆分，分隔行 |---| 跳过
    - 引用序号 [n] 原样保留（只是被上标，字符仍在）
    - 比对时压缩所有空白，容忍换行/run 拆分/首行缩进差异
    - docx 全文额外纳入页眉/页脚，以匹配封面被 —— 拆成主/副标题的完整标题

用法:
    python verify_fidelity.py <source.md> <output.docx> [--show N]
退出码: 0=内容一致(无缺失), 1=存在缺失, 2=参数错误
依赖: python-docx
"""
import re
import argparse
import sys
from docx import Document

LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
BOLD_RE = re.compile(r'\*\*(.+?)\*\*')
ITAL_RE = re.compile(r'\*(.+?)\*')
CODE_RE = re.compile(r'`(.+?)`')
SEP_CELL = re.compile(r'^:?-+:?$')
LIST_RE = re.compile(r'^(?:-|\d+\.)\s+(.*)$')


def strip_md(t):
    """去掉行内 markdown 标记，与转换器搬运结果对齐。"""
    t = LINK_RE.sub(r'\1', t)   # 链接 -> 仅文字
    t = BOLD_RE.sub(r'\1', t)   # 加粗
    t = ITAL_RE.sub(r'\1', t)   # 斜体
    t = CODE_RE.sub(r'\1', t)   # 代码
    return t


def squeeze(t):
    return re.sub(r'\s+', '', t)


def norm(t):
    return squeeze(strip_md(t))


def extract_source_units(path):
    """把源 md 拆成 (kind, text) 文本单元列表。"""
    lines = open(path, encoding='utf-8').read().split('\n')
    units = []
    i, n = 0, len(lines)
    in_fence = False
    while i < n:
        s = lines[i].strip()
        if s.startswith('```'):
            in_fence = not in_fence
            i += 1
            continue
        if in_fence:
            if s:
                units.append(('code', s))
            i += 1
            continue
        if s == '' or s == '---':
            i += 1
            continue
        if s.startswith('|'):                       # 表格块
            while i < n and lines[i].strip().startswith('|'):
                row = lines[i].strip()
                cells = [c.strip() for c in row.strip('|').split('|')]
                nonempty = [c for c in cells if c]
                if nonempty and all(SEP_CELL.match(c) for c in nonempty):
                    i += 1                           # 跳过分隔行
                    continue
                for c in cells:
                    if c:
                        units.append(('cell', c))
                i += 1
            continue
        if s.startswith('#'):                        # 标题
            lvl = len(s) - len(s.lstrip('#'))
            units.append(('head', s[lvl:].strip()))
            i += 1
            continue
        if s.startswith('>'):                        # 引用块
            units.append(('quote', re.sub(r'^>\s?', '', s)))
            i += 1
            continue
        m = LIST_RE.match(s)                          # 列表项
        if m:
            units.append(('list', m.group(1)))
            i += 1
            continue
        units.append(('para', s))                     # 普通段落
        i += 1
    return units


def extract_docx_text(doc):
    """docx 全文（段落 + 表格单元格 + 页眉页脚），用于正向子串匹配。"""
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            parts.append(p.text)
        for p in sec.footer.paragraphs:
            parts.append(p.text)
    return squeeze('\n'.join(parts))


# 反向抽查白名单：转换器自加的封面/目录/页眉页脚模板文本
BACK_WHITELIST = ('请在', '更新域', '目录', '目\u3000录', '日期：',
                  '第页/共', '页/共')


def main():
    ap = argparse.ArgumentParser(description='源 md ↔ docx 内容一致性审查')
    ap.add_argument('source', help='源 Markdown 路径')
    ap.add_argument('docx', help='生成的 .docx 路径')
    ap.add_argument('--show', type=int, default=20, help='最多展示的问题条数')
    args = ap.parse_args()

    units = extract_source_units(args.source)
    doc = Document(args.docx)
    dtext = extract_docx_text(doc)

    # ---- 正向覆盖：源每个单元必须出现在 docx ----
    by_kind = {}
    missing = []
    for kind, text in units:
        key = norm(text)
        by_kind.setdefault(kind, [0, 0])
        by_kind[kind][0] += 1
        if key and key in dtext:
            by_kind[kind][1] += 1
        elif key:
            # 拆分匹配：封面版本行在转换器中被按 "｜" 拆成多行居中显示（视觉优化，
            # 内容未丢）——整行不再连续存在于 docx，但逐段都存在即视为命中。
            if kind == 'quote' and re.search(r'[|｜]', key):
                segs = [s for s in re.split(r'[|｜]', key) if s]
                if segs and all(s in dtext for s in segs):
                    by_kind[kind][1] += 1
                else:
                    missing.append((kind, text))
            else:
                missing.append((kind, text))
        else:
            by_kind[kind][1] += 1  # 规范化后为空，视为无需比对

    total = sum(v[0] for v in by_kind.values())
    hit = sum(v[1] for v in by_kind.values())

    print('=' * 68)
    print('正向覆盖（源 -> docx，主判定）')
    print('=' * 68)
    for k in sorted(by_kind):
        tot, ok = by_kind[k]
        print('  %-6s 单元 %4d / 命中 %4d / 缺失 %d' % (k, tot, ok, tot - ok))
    cov = hit / total * 100 if total else 100.0
    print('  ------------------------------------------------')
    print('  合计 %d 个单元，命中 %d，缺失 %d，覆盖率 %.2f%%'
          % (total, hit, total - hit, cov))
    if missing:
        print('  缺失单元（前 %d 条）:' % args.show)
        for kind, text in missing[:args.show]:
            print('    [%s] %s' % (kind, text[:70]))

    # ---- 反向抽查：docx 正文疑似新增 ----
    src_full = norm(open(args.source, encoding='utf-8').read())
    suspicious = []
    for p in doc.paragraphs:
        raw = p.text.strip()
        key = squeeze(raw)
        if len(key) < 8:
            continue
        if any(w in raw for w in BACK_WHITELIST):
            continue
        if key in src_full:
            continue
        suspicious.append(raw)

    print('\n' + '=' * 68)
    print('反向抽查（docx -> 源，信息性；已排除封面/目录/页眉页脚模板）')
    print('=' * 68)
    print('  疑似不在源中的正文段落数 = %d' % len(suspicious))
    for s in suspicious[:args.show]:
        print('    ? %s' % s[:70])

    ok = (len(missing) == 0)
    print('\n结论:', '✅ 内容一致（源无遗漏）' if ok
          else '⚠ 存在 %d 处源内容未出现在 docx' % len(missing))
    sys.exit(0 if ok else 1)


if __name__ == '__main__':
    main()
